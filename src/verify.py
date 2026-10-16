"""Validate exported deliverables independently of the authoring engine."""
from pathlib import Path
import json
import math
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
import openpyxl
from pypdf import PdfReader

root=Path(__file__).resolve().parents[1]
summary=json.loads((root/'validation/summary.json').read_text(encoding='utf-8'))
book=root/'outputs/wildberries/Wildberries_аналитика.xlsx'
v=openpyxl.load_workbook(book,data_only=True)
f=openpyxl.load_workbook(book,data_only=False)
checks={}
expected={'Дашборд!C6':summary['revenue'],'Прогноз TOP10!E16':summary['top10_forecast'],'Запасы TOP20!L26':summary['top20_gap'],'Запасы TOP20!M26':summary['top20_excess'],'Проверка прогноза!L15':summary['wape'],'Проверка прогноза!M15':summary['naive_wape'],'Дашборд!L43':summary['low_conversion_count'],'Дашборд!L45':summary['marketing_count']}
for ref,target in expected.items():
    sheet,cell=ref.split('!'); actual=v[sheet][cell].value
    assert isinstance(actual,(int,float)) and math.isclose(actual,target,rel_tol=1e-10,abs_tol=1e-7),(ref,actual,target)
    checks[ref]=actual
assert len(v.sheetnames)==7
assert v['Ассортимент'].max_row>=2477
errors=[];formulas=0
for ws in f:
    for row in ws:
        for cell in row:
            formulas+=cell.data_type=='f'
for ws in v:
    for row in ws:
        for cell in row:
            if cell.data_type=='e':errors.append(f'{ws.title}!{cell.coordinate}: {cell.value}')
assert not errors,errors[:10]
assert f['Прогноз TOP10']['E5'].data_type=='f'
assert f['Запасы TOP20']['L5'].data_type=='f'
assert f['Ассортимент']['F5'].data_type=='f'
ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
with zipfile.ZipFile(book) as z:
    chartfiles=[x for x in z.namelist() if '/charts/chart' in x and x.endswith('.xml')]
    assert len(chartfiles)==2,chartfiles
    chart_refs={x:[el.text for el in ET.fromstring(z.read(x)).findall('.//c:f',ns)] for x in chartfiles}
    assert all(chart_refs.values()),chart_refs
    assert not any(x.startswith('xl/externalLinks/') for x in z.namelist())
    charts_xml='\n'.join(z.read(x).decode('utf-8') for x in chartfiles).upper()
    assert '853CC6' in charts_xml and 'E34C96' in charts_xml
pdf=root/'outputs/wildberries/Wildberries_отчёт.pdf'
reader=PdfReader(pdf)
assert len(reader.pages)==3
assert all(len(p.extract_text())>300 for p in reader.pages)
bt=pd.read_csv(root/'data/backtest_detail.csv')
assert len(bt)==90 and bt.status.eq('OK').all()
assert ((bt.prediction-(bt.older_orders+bt.prior_orders)/2).abs()<1e-8).all()
assert math.isclose(bt.abs_error.sum()/bt.actual.sum(),summary['wape'],abs_tol=1e-9)
result={'status':'PASS','sheets':v.sheetnames,'formula_count':formulas,'formula_errors':errors,'checks':checks,'charts':chart_refs,'pdf_pages':len(reader.pages),'backtest_observations':len(bt),'native_excel_tested':False}
(root/'validation/verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
